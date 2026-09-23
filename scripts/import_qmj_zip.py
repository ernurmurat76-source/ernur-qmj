#!/usr/bin/env python3
"""Convert a teacher-provided ZIP of DOCX lesson plans into server-side JSON references."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile
from xml.etree import ElementTree

from docx import Document


GROUPS = {
    "5кл Атамура Меруерт": (5, "Математика"),
    "6 класс атамұра Меруерт": (6, "Математика"),
    "7кл Атамура алг Меру": (7, "Алгебра"),
    "7кл Атамура геом Меру": (7, "Геометрия"),
    "8 класс алгебра атамұра Меруерт": (8, "Алгебра"),
    "8 класс геометрия атамұра Меруерт": (8, "Геометрия"),
    "9кл алг атамура Мер": (9, "Алгебра"),
    "9кл геом Атамура Меру": (9, "Геометрия"),
    "10 класс алг атамұра  Меру": (10, "Алгебра"),
    "10 класс гео атамұра Меру": (10, "Геометрия"),
    "11кл алгебра атамура ЖМБ": (11, "Алгебра"),
    "11кл геом атамура меру": (11, "Геометрия"),
}


def clean_line(value: str) -> str:
    value = value.replace("\uf0b7", "•").replace("\uf02d", "-")
    return re.sub(r"[ \t\u00a0]+", " ", value).strip()


def clean_text(value: str) -> str:
    lines = [clean_line(line) for line in (value or "").splitlines()]
    return "\n".join(line for line in lines if line)


def collapsed_cells(row) -> list[str]:
    result: list[str] = []
    for cell in row.cells:
        value = clean_text(cell.text)
        if not result or result[-1] != value:
            result.append(value)
    return result


def value_after_label(cells: list[str]) -> str:
    return next((value for value in reversed(cells[1:]) if value), "")


def lesson_numbers(filename: str) -> list[int]:
    stem = Path(filename).stem.lower()
    match = re.search(r"(\d{1,3})\s*[-–]\s*(\d{1,3})\s*(?:саб|сабақ)", stem)
    if match:
        start, end = map(int, match.groups())
        return list(range(start, end + 1)) if start <= end <= start + 5 else [start, end]
    match = re.search(r"(\d{1,3})\s*[-–]?\s*(?:саб|сабақ)", stem)
    return [int(match.group(1))] if match else []


def duration_minutes(label: str) -> int | None:
    match = re.search(r"(?<!\d)(\d{1,2})\s*(?:минут|мин)\b", label.lower())
    return int(match.group(1)) if match else None


def repair_broken_relationships(data: bytes) -> bytes:
    """Remove broken DOCX relationships (common Target=\"NULL\" export defect)."""
    source = BytesIO(data)
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as repaired:
        for info in original.infolist():
            payload = original.read(info)
            if info.filename.endswith(".rels"):
                try:
                    root = ElementTree.fromstring(payload)
                    for relationship in list(root):
                        value = relationship.attrib.get("Target", "").strip().lower()
                        if value in {"null", "/null", "word/null"}:
                            root.remove(relationship)
                    payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                except ElementTree.ParseError:
                    pass
            repaired.writestr(info, payload)
    return target.getvalue()


def load_document(data: bytes):
    try:
        return Document(BytesIO(data))
    except KeyError as error:
        if "NULL" not in str(error):
            raise
        return Document(BytesIO(repair_broken_relationships(data)))


def parse_document(data: bytes, group: str, filename: str) -> dict:
    grade, subject = GROUPS[group]
    document = load_document(data)
    if not document.tables:
        raise ValueError("Кесте табылмады")

    table = document.tables[0]
    meta = {"section": "", "topic": "", "objectives": "", "lesson_objectives": "", "values": ""}
    header_index = None

    for index, row in enumerate(table.rows):
        cells = collapsed_cells(row)
        joined = " ".join(cells).lower()
        label = cells[0].lower() if cells else ""
        if "уақыты" in joined and "педагог" in joined and "оқушы" in joined:
            header_index = index
            break
        value = value_after_label(cells)
        if label.startswith("бөлім"):
            meta["section"] = value
        elif "сабақтың тақырыбы" in label:
            meta["topic"] = value
        elif "оқу бағдарламасына сәйкес" in label and "мақсат" in label:
            meta["objectives"] = value
        elif label.startswith("сабақтың мақсаты"):
            meta["lesson_objectives"] = value
        elif "құндылық" in label:
            meta["values"] = value

    stages = []
    if header_index is not None:
        for row in table.rows[header_index + 1 :]:
            cells = collapsed_cells(row)
            if not any(cells):
                continue
            while cells and not cells[-1]:
                cells.pop()
            label = cells[0] if cells else ""
            stage = {
                "stage": label,
                "minutes": duration_minutes(label),
                "teacher": "",
                "learner": "",
                "assessment": "",
                "resources": "",
                "raw_columns": cells,
            }
            if len(cells) >= 5:
                stage.update(teacher=cells[1], learner=cells[2], assessment=cells[3], resources=cells[4])
            elif len(cells) == 4:
                stage.update(teacher=cells[1], assessment=cells[2], resources=cells[3])
            elif len(cells) == 3:
                stage.update(teacher=cells[1], learner=cells[2])
            elif len(cells) == 2:
                stage.update(teacher=cells[1])
            stages.append(stage)

    codes = re.findall(r"\b\d{1,2}(?:\.\d+){2,}\b", meta["objectives"])
    minutes = [stage["minutes"] for stage in stages if stage["minutes"] is not None]
    flags = []
    if not meta["topic"]:
        flags.append("missing_topic")
    if not meta["objectives"]:
        flags.append("missing_objectives")
    elif not codes:
        flags.append("objective_without_code")
    if not meta["lesson_objectives"]:
        flags.append("missing_lesson_objectives")
    if header_index is None:
        flags.append("missing_lesson_flow")
    if minutes and sum(minutes) != 45:
        flags.append("duration_not_45")

    source_key = f"{group}/{filename}"
    return {
        "id": hashlib.sha1(source_key.encode("utf-8")).hexdigest()[:16],
        "grade": grade,
        "subject": subject,
        "term": 1,
        "lesson_numbers": lesson_numbers(filename),
        "section": meta["section"],
        "topic": meta["topic"],
        "objectives": meta["objectives"],
        "objective_codes": codes,
        "lesson_objectives": meta["lesson_objectives"],
        "values": meta["values"],
        "stages": stages,
        "source": {"collection": "Атамұра 1 тоқсан Меруерт", "group": group, "file": filename},
        "quality_flags": flags,
    }


def main() -> int:
    if len(sys.argv) != 4:
        print("Usage: import_qmj_zip.py INPUT.zip OUTPUT.json REPORT.json", file=sys.stderr)
        return 2
    input_path, output_path, report_path = map(Path, sys.argv[1:])
    records = []
    failures = []

    with ZipFile(input_path) as archive:
        for info in archive.infolist():
            path = PurePosixPath(info.filename)
            if info.is_dir() or path.suffix.lower() != ".docx" or len(path.parts) < 3:
                continue
            group = path.parts[1]
            if group not in GROUPS:
                failures.append({"file": info.filename, "error": "Белгісіз топ"})
                continue
            try:
                records.append(parse_document(archive.read(info), group, path.name))
            except Exception as error:
                failures.append({"file": info.filename, "error": str(error)})

    records.sort(key=lambda item: (item["grade"], item["subject"], item["lesson_numbers"] or [999], item["topic"]))
    output = {
        "version": 1,
        "source_status": "teacher-provided",
        "source_scope": "1-тоқсан деп берілген архив; тек 5–11 сынып математикасы, алгебра және геометрия",
        "record_count": len(records),
        "records": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    flag_counts = Counter(flag for record in records for flag in record["quality_flags"])
    coverage = defaultdict(lambda: {"records": 0, "with_topic": 0, "with_objective_code": 0})
    for record in records:
        key = f'{record["grade"]}-сынып {record["subject"]}'
        coverage[key]["records"] += 1
        coverage[key]["with_topic"] += bool(record["topic"])
        coverage[key]["with_objective_code"] += bool(record["objective_codes"])
    report = {
        "records": len(records),
        "failures": failures,
        "quality_flags": dict(flag_counts),
        "coverage": dict(sorted(coverage.items())),
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
