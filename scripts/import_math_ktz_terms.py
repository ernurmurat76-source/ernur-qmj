#!/usr/bin/env python3
"""Import 2–4 term mathematics lessons from the teacher-provided 5–11 KTZ DOCX."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document


TABLES = [
    (5, "Математика", "Математика", ""),
    (6, "Математика", "Математика", ""),
    (7, "Алгебра", "Алгебра", ""),
    (7, "Геометрия", "Геометрия", ""),
    (8, "Алгебра", "Алгебра", ""),
    (8, "Геометрия", "Геометрия", ""),
    (9, "Алгебра", "Алгебра", ""),
    (9, "Геометрия", "Геометрия", ""),
    (10, "Алгебра", "Алгебра және анализ бастамалары", "ЖМБ"),
    (10, "Геометрия", "Геометрия", "ЖМБ"),
    (10, "Алгебра", "Алгебра және анализ бастамалары", "ҚГБ"),
    (10, "Геометрия", "Геометрия", "ҚГБ"),
    (11, "Алгебра", "Алгебра және анализ бастамалары", "ЖМБ"),
    (11, "Геометрия", "Геометрия", "ЖМБ"),
    (11, "Алгебра", "Алгебра және анализ бастамалары", "ҚГБ"),
    (11, "Геометрия", "Геометрия", "ҚГБ"),
]

VALUES = [
    "Тәуелсіздік және отаншылдық",
    "Бірлік және ынтымақтастық",
    "Әділдік және жауапкершілік",
    "Заң және тәртіп",
    "Еңбекқорлық және кәсіби біліктілік",
    "Жасампаздық және жаңашылдық",
]

GENERAL_GUIDANCE = [
    "Математикалық әрекетті нақты өмірлік жағдаятпен байланыстыру.",
    "Шешу тәсілін көрсету, нәтижені тексеру және қатемен жұмыс ұйымдастыру.",
    "Оқушыға бірнеше стратегияны салыстырып, таңдауын негіздету.",
    "Жұптық талқылау мен дербес математикалық әрекетті тең ұштастыру.",
]


def clean(value: str) -> str:
    value = (value or "").replace("\u00a0", " ").replace("\uf0b7", "•")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\s*\|\s*", "\n", value)
    value = re.sub(r"(?<=\d)(?=[А-ЯӘІҢҒҮҰҚӨҺа-яәіңғүұқөһ])", " ", value)
    lines = [line.strip(" -–;\t") for line in value.splitlines()]
    return "\n".join(line for line in lines if line)


def objective_codes(value: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"\b\d{1,2}(?:\.\d+){2,}\b", value)))


def find_columns(table) -> tuple[int, int, int]:
    headers = [clean(cell.text).lower() for cell in table.rows[0].cells]
    section = next((i for i, text in enumerate(headers) if "бөлім" in text), 1)
    topic = next((i for i, text in enumerate(headers) if "сабақ тақырыбы" in text), 2)
    objectives = next((i for i, text in enumerate(headers) if "оқу мақсат" in text), 3)
    return section, topic, objectives


def guidance_for(subject: str, grade: int) -> list[str]:
    extra = []
    if subject == "Геометрия":
        extra.append("Сызба, модель және дәлелдеуді байланыстырып, қажет жерде Қазақстанның сәулеті мен ою-өрнегін мәнмәтін ретінде қолдану.")
    else:
        extra.append("Есептеу, алгебралық түрлендіру немесе модельдеу нәтижесін кері амалмен не басқа тәсілмен тексерту.")
    if grade <= 9:
        extra.append("Контекстік есептің шартын схема, кесте немесе математикалық модель арқылы ұсыну.")
    else:
        extra.append("Күрделі есепте тиімді стратегияны таңдатып, математикалық негіздеме мен интерпретацияны талап ету.")
    return GENERAL_GUIDANCE + extra


def make_record(*, table_index: int, row_index: int, grade: int, subject: str,
                source_subject: str, track: str, term: int, lesson_number: int,
                section: str, topic: str, objectives: str, source_file: str) -> dict:
    key = f"{source_file}|{table_index}|{row_index}|{grade}|{subject}|{track}|{term}|{lesson_number}"
    flags = ["curriculum_only_no_reference_flow"]
    codes = objective_codes(objectives)
    if not topic:
        flags.append("missing_topic")
    if not objectives:
        flags.append("missing_objectives")
    elif not codes:
        flags.append("objective_without_code")
    return {
        "id": hashlib.sha1(key.encode("utf-8")).hexdigest()[:16],
        "grade": grade,
        "subject": subject,
        "source_subject": source_subject,
        "track": track,
        "term": term,
        "lesson_numbers": [lesson_number],
        "section": section,
        "topic": topic,
        "objectives": objectives,
        "objective_codes": codes,
        "lesson_objectives": "",
        "values": VALUES[(grade + term + lesson_number) % len(VALUES)],
        "methodological_guidance": guidance_for(subject, grade),
        "stages": [],
        "source": {
            "collection": "КТЖ 5–11 сынып Математика",
            "file": source_file,
            "table": table_index,
            "row": row_index,
            "status": "teacher-provided-ktz",
            "guidance": "Ы. Алтынсарин атындағы Ұлттық білім академиясының 2026–2027 әдістемелік нұсқау хаты",
        },
        "quality_flags": flags,
    }


def main() -> int:
    if len(sys.argv) != 5:
        print("Usage: import_math_ktz_terms.py KTZ.docx EXISTING.json OUTPUT.json REPORT.json", file=sys.stderr)
        return 2
    source_path, existing_path, output_path, report_path = map(Path, sys.argv[1:])
    document = Document(source_path)
    if len(document.tables) != len(TABLES):
        raise ValueError(f"Күтілген 16 кестенің орнына {len(document.tables)} кесте табылды")

    new_records = []
    for table_index, (table, metadata) in enumerate(zip(document.tables, TABLES)):
        grade, subject, source_subject, track = metadata
        section_column, topic_column, objective_column = find_columns(table)
        term = None
        current_section = ""
        for row_index, row in enumerate(table.rows[1:], start=1):
            cells = [clean(cell.text) for cell in row.cells]
            first = cells[0] if cells else ""
            unique_text = " ".join(dict.fromkeys(value for value in cells if value))
            term_match = re.search(r"\b([1-4])\s*[-–]?\s*тоқсан\b", unique_text, re.I)
            if term_match and not re.fullmatch(r"\d+", first):
                term = int(term_match.group(1))
                continue
            if term not in {2, 3, 4} or not re.fullmatch(r"\d+", first):
                continue
            section_value = cells[section_column] if section_column < len(cells) else ""
            if section_value:
                current_section = section_value
            topic = cells[topic_column] if topic_column < len(cells) else ""
            objectives = cells[objective_column] if objective_column < len(cells) else ""
            if not topic and not objectives:
                continue
            new_records.append(make_record(
                table_index=table_index, row_index=row_index, grade=grade, subject=subject,
                source_subject=source_subject, track=track, term=term, lesson_number=int(first),
                section=current_section, topic=topic, objectives=objectives, source_file=source_path.name,
            ))

    existing = json.loads(existing_path.read_text(encoding="utf-8"))
    retained = [record for record in existing.get("records", []) if not (
        record.get("source", {}).get("status") == "teacher-provided-ktz" and record.get("term") in {2, 3, 4}
    )]
    records = retained + new_records
    records.sort(key=lambda item: (
        item.get("grade", 0), item.get("subject", ""), item.get("track", ""), item.get("term", 0),
        (item.get("lesson_numbers") or [999])[0], item.get("topic", "")
    ))
    output = {
        "version": 2,
        "source_status": "mixed-teacher-provided",
        "source_scope": "1-тоқсан: мұғалім ҚМЖ архиві; 2–4-тоқсан: мұғалім берген 5–11 сынып математика КТЖ; әдістемелік бағыт: ҰБА 2026–2027 нұсқау хаты",
        "record_count": len(records),
        "records": records,
    }
    output_path.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    coverage = defaultdict(lambda: {"records": 0, "with_codes": 0, "missing_objectives": 0})
    for record in new_records:
        key = f'{record["grade"]}-сынып {record["subject"]} {record["track"] or "жалпы"} {record["term"]}-тоқсан'
        coverage[key]["records"] += 1
        coverage[key]["with_codes"] += bool(record["objective_codes"])
        coverage[key]["missing_objectives"] += not bool(record["objectives"])
    flags = Counter(flag for record in new_records for flag in record["quality_flags"])
    report = {
        "source_file": source_path.name,
        "existing_records_retained": len(retained),
        "new_records": len(new_records),
        "total_records": len(records),
        "quality_flags": dict(flags),
        "coverage": dict(sorted(coverage.items())),
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
